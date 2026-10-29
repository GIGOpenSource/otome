"""一部游戏的配置表。用素材名称和剧情编号互相引用，不写数据库 ID，导出的表格可以直接再导入。"""
import json
import logging
from pathlib import Path

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from Apps.views.Otome.content_i18n import (
    ASSET_FIELDS,
    ASSET_LIMITS,
    CHARACTER_FIELDS,
    CHARACTER_LIMITS,
    CHOICE_FIELDS,
    CHOICE_LIMITS,
    GALLERY_FIELDS,
    GALLERY_LIMITS,
    GAME_FIELDS,
    GAME_LIMITS,
    NODE_FIELDS,
    NODE_LIMITS,
    clean_texts,
)
from Apps.views.Otome.inserts import sync_insert_gallery
from Apps.views.Otome.landing import clean_ios, clean_package
from models.models import (
    OtomeAsset,
    OtomeCharacter,
    OtomeChoice,
    OtomeGalleryItem,
    OtomeGame,
    OtomeCharacterImage,
    OtomeNode,
    OtomeNodeFigure,
)

logger = logging.getLogger(__name__)

SHEET_GAME = "游戏"
SHEET_ASSET = "素材"
SHEET_CHARACTER = "角色"
SHEET_NODE = "剧情"
SHEET_CHOICE = "选项"
SHEET_GALLERY = "解锁"
SHEET_HELP = "填写说明"

TYPE_TO_LABEL = {
    OtomeAsset.TYPE_IMAGE: "图片",
    OtomeAsset.TYPE_AUDIO: "音频",
    OtomeAsset.TYPE_VIDEO: "视频",
}
LABEL_TO_TYPE = {label: key for key, label in TYPE_TO_LABEL.items()}
USAGE_TO_LABEL = {"": "", "background": "背景", "character": "立绘", "cover": "封面"}
LABEL_TO_USAGE = {label: key for key, label in USAGE_TO_LABEL.items() if label}
SCREEN_TO_LABEL = {"": "", "portrait": "竖屏", "landscape": "横屏"}
LABEL_TO_SCREEN = {label: key for key, label in SCREEN_TO_LABEL.items() if label}
STATUS_TO_LABEL = {OtomeGame.STATUS_DRAFT: "草稿", OtomeGame.STATUS_PUBLISHED: "已发布"}
LABEL_TO_STATUS = {label: key for key, label in STATUS_TO_LABEL.items()}
KIND_TO_LABEL = {OtomeGalleryItem.KIND_VIDEO: "视频", OtomeGalleryItem.KIND_PORTRAIT: "立绘"}
LABEL_TO_KIND = {label: key for key, label in KIND_TO_LABEL.items()}

GAME_HEADERS = ["故事编码", "标题", "副标题", "简介", "竖屏封面", "横屏封面", "Google包名", "App Store ID", "状态", "排序"]
ASSET_HEADERS = ["名称", "类型", "用途", "屏幕方向", "外链地址", "文件名", "备注"]
CHARACTER_HEADERS = ["姓名", "身份", "竖屏立绘", "横屏立绘", "名字颜色", "介绍", "排序"]
NODE_HEADERS = [
    "剧情编号", "排序", "说话角色", "显示名", "对白",
    "竖屏背景", "横屏背景", "竖屏立绘", "横屏立绘",
    "背景音乐", "语音", "交互音频", "结束后插入", "解锁标题", "下一句",
    "开场", "关键节点", "结局", "结局标题",
]
CHOICE_HEADERS = ["所在剧情", "排序", "选项文字", "跳转剧情", "交互音频"]
GALLERY_HEADERS = ["类型", "标题", "素材", "解锁剧情", "默认解锁", "排序", "来源"]
TEXT_COLUMN = "多语言"

HELP_ROWS = [
    ("游戏", "只填一行。故事编码用来识别是不是同一部游戏，再次导入会更新这一部，不会新建重复的。状态填「草稿」或「已发布」。封面填素材表里的名称。"),
    ("素材", "类型填「图片」「音频」或「视频」。图片的用途填「背景」「立绘」或「封面」，屏幕方向填「竖屏」或「横屏」。有上传文件时外链留空，导入会沿用原来的文件。只有外链时填写外链地址。"),
    ("角色", "竖屏立绘、横屏立绘填素材名称。名字颜色用 #e7c27a 这种写法。"),
    ("剧情", "剧情编号在这一部游戏里不能重复。说话角色填角色姓名，旁白就留空。背景、立绘、音乐、语音都填素材名称。下一句填剧情编号。开场、关键节点、结局填「是」或「否」。一部游戏只标一个开场。"),
    ("选项", "所在剧情和跳转剧情都填剧情编号。有选项时，玩家点选项跳转，不再走「下一句」。"),
    ("解锁", "类型填「立绘」或「视频」。来源填「手动」。来源是「剧情插入」的行由剧情表的「结束后插入」生成，导入时会跳过，改剧情表即可。"),
    ("引用", "表格里不出现数据库 ID。素材靠名称对应，剧情靠剧情编号对应，角色靠姓名对应。同一张表里这些名称不要重复。"),
    ("多语言", "可选列。主语言仍写在原来的标题、对白等列里。其他语言写成 JSON，例如 {\"en\":{\"title\":\"Hello\"}}。这一列留空会清掉已有译文；旧表格没有这一列时，导入不会改动译文。"),
]


class SheetError(Exception):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("\n".join(errors))


def _yes(value):
    return "是" if value else "否"


def _asset_name(asset):
    if asset is None or asset.is_deleted:
        return ""
    return asset.name


def _file_name(asset):
    if not asset.file:
        return ""
    return Path(asset.file.name).name


def export_workbook(game):
    book = Workbook()
    _write_help(book.active)
    _write_sheet(book, SHEET_GAME, GAME_HEADERS, [_game_row(game)])
    assets = list(game.assets.filter(is_deleted=False).order_by("id"))
    _write_sheet(book, SHEET_ASSET, ASSET_HEADERS, [_asset_row(item) for item in assets])
    characters = list(game.characters.filter(is_deleted=False).order_by("sort_order", "id"))
    _write_sheet(book, SHEET_CHARACTER, CHARACTER_HEADERS, [_character_row(item) for item in characters])
    nodes = list(
        game.nodes.filter(is_deleted=False)
        .select_related(
            "character", "background", "background_land", "portrait", "portrait_land",
            "bgm", "voice", "interact", "insert_media", "next_node",
        )
        .order_by("sort_order", "id")
    )
    _write_sheet(book, SHEET_NODE, NODE_HEADERS, [_node_row(item) for item in nodes])
    choices = (
        OtomeChoice.objects.filter(node__in=nodes)
        .select_related("node", "next_node", "interact")
        .order_by("node__sort_order", "node_id", "sort_order", "id")
    )
    _write_sheet(book, SHEET_CHOICE, CHOICE_HEADERS, [_choice_row(item) for item in choices])
    gallery = list(game.gallery_items.filter(is_deleted=False).select_related("asset", "unlock_node").order_by("sort_order", "id"))
    _write_sheet(book, SHEET_GALLERY, GALLERY_HEADERS, [_gallery_row(item) for item in gallery])
    return book


def export_game(game, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    export_workbook(game).save(path)
    logger.info("导出游戏配置表 game=%s path=%s", game.id, path)
    return path


def _ensure_sheet_figure(node):
    """表格只有一张立绘。已经在编辑器里摆过多名角色时，不覆盖。"""
    if node.figures.exists() or not node.portrait_id:
        return
    image = OtomeCharacterImage.objects.filter(
        image_id=node.portrait_id, character__game_id=node.game_id, character__is_deleted=False
    ).first()
    OtomeNodeFigure.objects.create(
        node=node,
        character=image.character if image else node.character,
        image=image,
        position="e",
        sort_order=0,
    )


def import_game(path, dry_run=False):
    path = Path(path)
    if not path.is_file():
        raise SheetError([f"找不到表格：{path}"])
    book = load_workbook(path, data_only=True)
    payload, errors = _read_book(book)
    if errors:
        raise SheetError(errors)
    if dry_run:
        return {"dry_run": True, **_summary(payload)}
    with transaction.atomic():
        game = _apply(payload)
    logger.info("导入游戏配置表 game=%s code=%s", game.id, game.code)
    return {"dry_run": False, "id": game.id, "code": game.code, "title": game.title, **_summary(payload)}


def _summary(payload):
    return {
        "title": payload["game"]["标题"],
        "assets": len(payload["assets"]),
        "characters": len(payload["characters"]),
        "nodes": len(payload["nodes"]),
        "choices": len(payload["choices"]),
        "gallery": len(payload["gallery"]),
    }


def _game_row(game):
    return [
        game.code or "",
        game.title,
        game.subtitle,
        game.summary,
        _asset_name(game.cover),
        _asset_name(game.cover_land),
        game.play_package,
        game.ios_app_id,
        STATUS_TO_LABEL.get(game.status, "草稿"),
        game.sort_order,
        _texts_cell(game),
    ]


def _asset_row(item):
    return [
        item.name,
        TYPE_TO_LABEL.get(item.asset_type, item.asset_type),
        USAGE_TO_LABEL.get(item.usage, ""),
        SCREEN_TO_LABEL.get(item.screen, ""),
        item.url,
        _file_name(item),
        item.remark,
        _texts_cell(item),
    ]


def _character_row(item):
    return [
        item.name,
        item.title,
        _asset_name(item.portrait),
        _asset_name(item.portrait_land),
        item.name_color,
        item.intro,
        item.sort_order,
        _texts_cell(item),
    ]


def _node_row(item):
    return [
        item.node_key,
        item.sort_order,
        item.character.name if item.character_id and item.character and not item.character.is_deleted else "",
        item.speaker_name,
        item.content,
        _asset_name(item.background),
        _asset_name(item.background_land),
        _asset_name(item.portrait),
        _asset_name(item.portrait_land),
        _asset_name(item.bgm),
        _asset_name(item.voice),
        _asset_name(item.interact),
        _asset_name(item.insert_media),
        item.insert_title,
        item.next_node.node_key if item.next_node_id and item.next_node and not item.next_node.is_deleted else "",
        _yes(item.is_start),
        _yes(item.is_key),
        _yes(item.is_ending),
        item.ending_title,
        _texts_cell(item),
    ]


def _choice_row(item):
    return [
        item.node.node_key,
        item.sort_order,
        item.text,
        item.next_node.node_key if item.next_node_id and not item.next_node.is_deleted else "",
        _asset_name(item.interact),
        _texts_cell(item),
    ]


def _gallery_row(item):
    return [
        KIND_TO_LABEL.get(item.kind, item.kind),
        item.title,
        _asset_name(item.asset),
        item.unlock_node.node_key if item.unlock_node_id and item.unlock_node and not item.unlock_node.is_deleted else "",
        _yes(item.unlocked_by_default),
        item.sort_order,
        "剧情插入" if item.is_insert else "手动",
        _texts_cell(item),
    ]


def _write_help(sheet):
    sheet.title = SHEET_HELP
    _write_table(sheet, ["表", "怎么填"], HELP_ROWS)


def _write_sheet(book, title, headers, rows):
    _write_table(book.create_sheet(title), [*headers, TEXT_COLUMN], rows)


def _write_table(sheet, headers, rows):
    header_fill = PatternFill("solid", fgColor="241C16")
    header_font = Font(color="F6EFE4", bold=True)
    wrap = Alignment(wrap_text=True, vertical="center")
    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(1, col, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    for row_index, row in enumerate(rows, start=2):
        for col, value in enumerate(row, start=1):
            cell = sheet.cell(row_index, col, value)
            cell.alignment = wrap
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(len(rows) + 1, 1)}"
    for col, header in enumerate(headers, start=1):
        width = 18
        if header in ("对白", "简介", "介绍", "怎么填"):
            width = 42
        elif header in ("名称", "标题", "故事编码", TEXT_COLUMN):
            width = 28 if header == TEXT_COLUMN else 22
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.row_dimensions[1].height = 22


def _read_book(book):
    errors = []
    if SHEET_GAME not in book.sheetnames:
        return None, ["缺少「游戏」表"]
    games = _read_sheet(book[SHEET_GAME], GAME_HEADERS, errors)
    data_rows = [row for row in games if _any_value(row)]
    if len(data_rows) != 1:
        errors.append("「游戏」表只能有一行，一份表格对应一部游戏")
        return None, errors
    game = data_rows[0]
    code = _text(game.get("故事编码"))
    title = _text(game.get("标题"))
    line = game.get("_row") or 2
    if not code:
        errors.append(f"「游戏」第 {line} 行：故事编码不能为空，导入靠它识别这一部")
    if not title:
        errors.append(f"「游戏」第 {line} 行：标题不能为空")
    status = _text(game.get("状态")) or "草稿"
    if status not in LABEL_TO_STATUS:
        errors.append(f"「游戏」第 {line} 行：状态只能填「草稿」或「已发布」")
    package = clean_package(game.get("Google包名"))
    if package is None:
        errors.append(f"「游戏」第 {line} 行：Google 包名格式不正确")
    ios_id = clean_ios(game.get("App Store ID"))
    if ios_id is None:
        errors.append(f"「游戏」第 {line} 行：App Store ID 应是数字")
    game["故事编码"] = code
    game["标题"] = title
    game["状态"] = status
    game["Google包名"] = package or ""
    game["App Store ID"] = ios_id or ""
    game["排序"] = _number(game.get("排序"), "游戏", 2, "排序", errors)

    assets = _read_sheet(book[SHEET_ASSET], ASSET_HEADERS, errors) if SHEET_ASSET in book.sheetnames else []
    characters = _read_sheet(book[SHEET_CHARACTER], CHARACTER_HEADERS, errors) if SHEET_CHARACTER in book.sheetnames else []
    nodes = _read_sheet(book[SHEET_NODE], NODE_HEADERS, errors) if SHEET_NODE in book.sheetnames else []
    choices = _read_sheet(book[SHEET_CHOICE], CHOICE_HEADERS, errors) if SHEET_CHOICE in book.sheetnames else []
    gallery = _read_sheet(book[SHEET_GALLERY], GALLERY_HEADERS, errors) if SHEET_GALLERY in book.sheetnames else []
    assets = [row for row in assets if _any_value(row)]
    characters = [row for row in characters if _any_value(row)]
    nodes = [row for row in nodes if _any_value(row)]
    choices = [row for row in choices if _any_value(row)]
    gallery = [row for row in gallery if _any_value(row)]

    asset_names = _unique_names(assets, "名称", SHEET_ASSET, errors)
    character_names = _unique_names(characters, "姓名", SHEET_CHARACTER, errors)
    node_keys = _unique_names(nodes, "剧情编号", SHEET_NODE, errors)
    _check_assets(assets, errors)
    for field in ("竖屏封面", "横屏封面"):
        cover_name = _text(game.get(field))
        game[field] = cover_name
        if cover_name and cover_name not in asset_names:
            errors.append(f"「游戏」第 {line} 行：{field}「{cover_name}」不在素材表里")
        elif cover_name and asset_names.get(cover_name) not in ("图片", "视频"):
            errors.append(f"「游戏」第 {line} 行：{field}「{cover_name}」必须是图片或视频")
    _check_characters(characters, asset_names, errors)
    _check_nodes(nodes, character_names, asset_names, node_keys, errors)
    _check_choices(choices, node_keys, asset_names, errors)
    manual_gallery = _check_gallery(gallery, asset_names, node_keys, errors)
    if errors:
        return None, errors
    return {
        "game": game,
        "assets": assets,
        "characters": characters,
        "nodes": nodes,
        "choices": choices,
        "gallery": manual_gallery,
    }, []


def _read_sheet(sheet, headers, errors):
    header_row = next(sheet.iter_rows(min_row=1, max_row=1), None)
    if header_row is None:
        errors.append(f"「{sheet.title}」表是空的")
        return []
    found = [_text(cell.value) for cell in header_row]
    missing = [header for header in headers if header not in found]
    if missing:
        errors.append(f"「{sheet.title}」表缺少列：{'、'.join(missing)}")
        return []
    index = {name: pos for pos, name in enumerate(found)}
    text_pos = index.get(TEXT_COLUMN)
    rows = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if values is None:
            continue
        row = {"_row": excel_row}
        for header in headers:
            pos = index[header]
            row[header] = values[pos] if pos < len(values) else ""
        if text_pos is not None:
            row[TEXT_COLUMN] = values[text_pos] if values and text_pos < len(values) else ""
        rows.append(row)
    return rows


def _check_assets(rows, errors):
    for row in rows:
        line = row["_row"]
        name = _text(row.get("名称"))
        if not name:
            errors.append(f"「素材」第 {line} 行：名称不能为空")
            continue
        row["名称"] = name
        kind = _text(row.get("类型"))
        if kind not in LABEL_TO_TYPE:
            errors.append(f"「素材」第 {line} 行：类型只能填图片、音频或视频")
            continue
        row["类型"] = kind
        usage = _text(row.get("用途"))
        screen = _text(row.get("屏幕方向"))
        if kind == "图片" and (usage or screen):
            if usage not in LABEL_TO_USAGE:
                errors.append(f"「素材」第 {line} 行：用途只能填背景、立绘或封面")
            if screen not in LABEL_TO_SCREEN:
                errors.append(f"「素材」第 {line} 行：屏幕方向只能填竖屏或横屏")
        elif kind != "图片" and (usage or screen):
            errors.append(f"「素材」第 {line} 行：音频和视频不用填用途和屏幕方向")
        row["用途"] = usage
        row["屏幕方向"] = screen
        row["外链地址"] = _text(row.get("外链地址"))
        row["备注"] = _text(row.get("备注"))


def _check_characters(rows, asset_names, errors):
    for row in rows:
        line = row["_row"]
        name = _text(row.get("姓名"))
        if not name:
            errors.append(f"「角色」第 {line} 行：姓名不能为空")
        row["姓名"] = name
        row["身份"] = _text(row.get("身份"))
        row["名字颜色"] = _text(row.get("名字颜色")) or "#e7c27a"
        row["介绍"] = _text(row.get("介绍"))
        row["排序"] = _number(row.get("排序"), "角色", line, "排序", errors)
        _need_asset(row, "竖屏立绘", asset_names, "角色", line, errors, visual=True)
        _need_asset(row, "横屏立绘", asset_names, "角色", line, errors, visual=True)


def _check_nodes(rows, character_names, asset_names, node_keys, errors):
    starts = 0
    visual_fields = ("竖屏背景", "横屏背景", "竖屏立绘", "横屏立绘", "结束后插入")
    audio_fields = ("背景音乐", "语音", "交互音频")
    for row in rows:
        line = row["_row"]
        key = _text(row.get("剧情编号"))
        if not key:
            errors.append(f"「剧情」第 {line} 行：剧情编号不能为空")
        row["剧情编号"] = key
        row["排序"] = _number(row.get("排序"), "剧情", line, "排序", errors)
        speaker = _text(row.get("说话角色"))
        if speaker and speaker not in character_names:
            errors.append(f"「剧情」第 {line} 行：说话角色「{speaker}」不在角色表里")
        row["说话角色"] = speaker
        row["显示名"] = _text(row.get("显示名"))
        row["对白"] = _text(row.get("对白"))
        if not row["对白"]:
            errors.append(f"「剧情」第 {line} 行：对白不能为空")
        for field in visual_fields:
            _need_asset(row, field, asset_names, "剧情", line, errors, visual=True)
        for field in audio_fields:
            _need_asset(row, field, asset_names, "剧情", line, errors, audio=True)
        row["解锁标题"] = _text(row.get("解锁标题"))
        nxt = _text(row.get("下一句"))
        if nxt and nxt not in node_keys:
            errors.append(f"「剧情」第 {line} 行：下一句「{nxt}」不在剧情表里")
        row["下一句"] = nxt
        row["开场"] = _bool_cell(row.get("开场"), "剧情", line, "开场", errors)
        row["关键节点"] = _bool_cell(row.get("关键节点"), "剧情", line, "关键节点", errors)
        row["结局"] = _bool_cell(row.get("结局"), "剧情", line, "结局", errors)
        row["结局标题"] = _text(row.get("结局标题"))
        if row["开场"]:
            starts += 1
    if starts > 1:
        errors.append("「剧情」表里只能有一行开场填「是」")


def _check_choices(rows, node_keys, asset_names, errors):
    for row in rows:
        line = row["_row"]
        node_key = _text(row.get("所在剧情"))
        target = _text(row.get("跳转剧情"))
        text = _text(row.get("选项文字"))
        if node_key not in node_keys:
            errors.append(f"「选项」第 {line} 行：所在剧情「{node_key}」不在剧情表里")
        if not text:
            errors.append(f"「选项」第 {line} 行：选项文字不能为空")
        if target not in node_keys:
            errors.append(f"「选项」第 {line} 行：跳转剧情「{target}」不在剧情表里")
        row["所在剧情"] = node_key
        row["跳转剧情"] = target
        row["选项文字"] = text
        row["排序"] = _number(row.get("排序"), "选项", line, "排序", errors)
        _need_asset(row, "交互音频", asset_names, "选项", line, errors, audio=True)


def _check_gallery(rows, asset_names, node_keys, errors):
    manual = []
    for row in rows:
        line = row["_row"]
        source = _text(row.get("来源")) or "手动"
        if source == "剧情插入":
            continue
        if source != "手动":
            errors.append(f"「解锁」第 {line} 行：来源只能填「手动」或「剧情插入」")
        kind = _text(row.get("类型"))
        if kind not in LABEL_TO_KIND:
            errors.append(f"「解锁」第 {line} 行：类型只能填立绘或视频")
        title = _text(row.get("标题"))
        if not title:
            errors.append(f"「解锁」第 {line} 行：标题不能为空")
        _need_asset(row, "素材", asset_names, "解锁", line, errors)
        unlock = _text(row.get("解锁剧情"))
        if unlock and unlock not in node_keys:
            errors.append(f"「解锁」第 {line} 行：解锁剧情「{unlock}」不在剧情表里")
        row["来源"] = "手动"
        row["类型"] = kind
        row["标题"] = title
        row["解锁剧情"] = unlock
        row["默认解锁"] = _bool_cell(row.get("默认解锁"), "解锁", line, "默认解锁", errors)
        row["排序"] = _number(row.get("排序"), "解锁", line, "排序", errors)
        manual.append(row)
    return manual


def _need_asset(row, field, names, sheet, line, errors, visual=False, audio=False):
    name = _text(row.get(field))
    row[field] = name
    if not name:
        return
    if name not in names:
        errors.append(f"「{sheet}」第 {line} 行：{field}「{name}」不在素材表里")
        return
    kind = names[name]
    if visual and kind not in ("图片", "视频"):
        errors.append(f"「{sheet}」第 {line} 行：{field}「{name}」必须是图片或视频")
    if audio and kind != "音频":
        errors.append(f"「{sheet}」第 {line} 行：{field}「{name}」必须是音频")


def _unique_names(rows, field, sheet, errors):
    found = {}
    for row in rows:
        name = _text(row.get(field))
        if not name:
            continue
        if name in found:
            errors.append(f"「{sheet}」第 {row['_row']} 行：{field}「{name}」和第 {found[name]} 行重复")
            continue
        found[name] = row["_row"]
        if field == "名称":
            found[name] = _text(row.get("类型")) or found[name]
    if field == "名称":
        names = {}
        seen = {}
        for row in rows:
            name = _text(row.get("名称"))
            if not name:
                continue
            if name in seen:
                continue
            seen[name] = row["_row"]
            names[name] = _text(row.get("类型"))
        return names
    return set(found)


def _apply(payload):
    game_row = payload["game"]
    code = game_row["故事编码"]
    game = OtomeGame.objects.filter(code=code).first() or OtomeGame(code=code)
    game.title = game_row["标题"][:100]
    game.subtitle = _text(game_row.get("副标题"))[:200]
    game.summary = _text(game_row.get("简介"))
    game.play_package = game_row["Google包名"]
    game.ios_app_id = game_row["App Store ID"]
    game.status = LABEL_TO_STATUS[game_row["状态"]]
    game.sort_order = game_row["排序"]
    game.is_deleted = False
    _assign_texts(game, game_row, GAME_FIELDS, GAME_LIMITS)
    game.save()

    assets = {}
    for row in payload["assets"]:
        asset = OtomeAsset.objects.filter(game=game, name=row["名称"], is_deleted=False).first()
        if asset is None:
            taken = OtomeAsset.objects.filter(name=row["名称"], is_deleted=False).exclude(game=game).exclude(game__isnull=True).first()
            if taken is not None:
                raise SheetError([f"素材「{row['名称']}」已经属于别的游戏"])
            asset = OtomeAsset.objects.filter(name=row["名称"], game__isnull=True, is_deleted=False).first() or OtomeAsset(name=row["名称"])
        asset.game = game
        asset.name = row["名称"][:100]
        asset.asset_type = LABEL_TO_TYPE[row["类型"]]
        asset.usage = LABEL_TO_USAGE.get(row["用途"], "") if asset.asset_type == OtomeAsset.TYPE_IMAGE else ""
        asset.screen = LABEL_TO_SCREEN.get(row["屏幕方向"], "") if asset.asset_type == OtomeAsset.TYPE_IMAGE else ""
        asset.remark = row["备注"][:200]
        asset.is_deleted = False
        _assign_texts(asset, row, ASSET_FIELDS, ASSET_LIMITS)
        if row["外链地址"]:
            if asset.file:
                asset.file.delete(save=False)
            asset.file = ""
            asset.url = row["外链地址"][:1000]
        asset.save()
        assets[asset.name] = asset

    characters = {}
    kept_names = []
    for row in payload["characters"]:
        character = OtomeCharacter.objects.filter(game=game, name=row["姓名"], is_deleted=False).first() or OtomeCharacter(game=game, name=row["姓名"])
        character.name = row["姓名"][:50]
        character.title = row["身份"][:100]
        character.portrait = assets.get(row["竖屏立绘"])
        character.portrait_land = assets.get(row["横屏立绘"])
        character.name_color = row["名字颜色"][:20]
        character.intro = row["介绍"]
        character.sort_order = row["排序"]
        character.is_deleted = False
        _assign_texts(character, row, CHARACTER_FIELDS, CHARACTER_LIMITS)
        character.save()
        characters[character.name] = character
        kept_names.append(character.name)
    game.characters.filter(is_deleted=False).exclude(name__in=kept_names).update(is_deleted=True)

    nodes = {}
    kept_keys = []
    for row in payload["nodes"]:
        node = OtomeNode.objects.filter(game=game, node_key=row["剧情编号"], is_deleted=False).first() or OtomeNode(game=game, node_key=row["剧情编号"])
        node.node_key = row["剧情编号"][:64]
        node.sort_order = row["排序"]
        node.character = characters.get(row["说话角色"])
        node.speaker_name = row["显示名"][:50]
        node.content = row["对白"]
        node.background = assets.get(row["竖屏背景"])
        node.background_land = assets.get(row["横屏背景"])
        node.portrait = assets.get(row["竖屏立绘"])
        node.portrait_land = assets.get(row["横屏立绘"])
        node.bgm = assets.get(row["背景音乐"])
        node.voice = assets.get(row["语音"])
        node.interact = assets.get(row["交互音频"])
        node.insert_media = assets.get(row["结束后插入"])
        node.insert_title = row["解锁标题"][:100]
        node.is_start = row["开场"]
        node.is_key = row["关键节点"]
        node.is_ending = row["结局"]
        node.ending_title = row["结局标题"][:100]
        node.is_deleted = False
        _assign_texts(node, row, NODE_FIELDS, NODE_LIMITS)
        node.next_node = None
        node.save()
        _ensure_sheet_figure(node)
        nodes[node.node_key] = node
        kept_keys.append(node.node_key)
    for row in payload["nodes"]:
        node = nodes[row["剧情编号"]]
        node.next_node = nodes.get(row["下一句"])
        node.save(update_fields=["next_node", "updated_at"])
    OtomeChoice.objects.filter(node__game=game).delete()
    for row in payload["choices"]:
        choice = OtomeChoice(
            node=nodes[row["所在剧情"]],
            text=row["选项文字"][:200],
            next_node=nodes[row["跳转剧情"]],
            interact=assets.get(row["交互音频"]),
            sort_order=row["排序"],
        )
        _assign_texts(choice, row, CHOICE_FIELDS, CHOICE_LIMITS, choice.node.lang)
        choice.save()
    stale = list(game.nodes.filter(is_deleted=False).exclude(node_key__in=kept_keys))
    for node in stale:
        node.is_deleted = True
        node.next_node = None
        node.is_start = False
        node.save(update_fields=["is_deleted", "next_node", "is_start", "updated_at"])

    game.cover = assets.get(_text(game_row.get("竖屏封面")))
    game.cover_land = assets.get(_text(game_row.get("横屏封面")))
    game.save(update_fields=["cover", "cover_land", "updated_at"])

    game.gallery_items.filter(is_insert=False).update(is_deleted=True)
    for row in payload["gallery"]:
        kind = LABEL_TO_KIND[row["类型"]]
        item = OtomeGalleryItem.objects.filter(game=game, kind=kind, title=row["标题"], is_insert=False).first() or OtomeGalleryItem(game=game, is_insert=False)
        item.kind = kind
        item.title = row["标题"][:100]
        item.asset = assets.get(row["素材"])
        item.unlock_node = nodes.get(row["解锁剧情"])
        item.unlocked_by_default = row["默认解锁"]
        item.sort_order = row["排序"]
        item.is_deleted = False
        _assign_texts(item, row, GALLERY_FIELDS, GALLERY_LIMITS)
        item.save()
    for node in nodes.values():
        sync_insert_gallery(node)
    return game


def _any_value(row):
    return any(_text(value) for key, value in row.items() if key != "_row")


def _texts_cell(item):
    raw = getattr(item, "texts", None) or {}
    if not raw:
        return ""
    return json.dumps(raw, ensure_ascii=False, separators=(",", ":"))


def _assign_texts(obj, row, fields, limits, source=None):
    if TEXT_COLUMN not in row:
        return
    source = source or getattr(obj, "lang", None) or "zh-Hans"
    line = row.get("_row") or ""
    try:
        obj.texts = clean_texts(row.get(TEXT_COLUMN), fields, limits, source)
    except ValueError as exc:
        raise SheetError([f"第 {line} 行：多语言需要是 JSON，例如 {{\"en\":{{\"title\":\"Hello\"}}}}"]) from exc


def _text(value):
    if value is None:
        return ""
    return str(value).strip()


def _number(value, sheet, line, label, errors):
    text = _text(value)
    if not text:
        return 0
    try:
        return int(float(text))
    except ValueError:
        errors.append(f"「{sheet}」第 {line} 行：{label}必须是数字")
        return 0


def _bool_cell(value, sheet, line, label, errors):
    text = _text(value)
    if text in ("", "否", "0", "false", "False", "no", "No"):
        return False
    if text in ("是", "1", "true", "True", "yes", "Yes"):
        return True
    errors.append(f"「{sheet}」第 {line} 行：{label}只能填「是」或「否」")
    return False
